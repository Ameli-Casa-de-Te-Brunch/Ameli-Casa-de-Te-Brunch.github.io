#!/usr/bin/env python3
"""Lee desde una pestaña de Google Sheets "publicada en la web" (CSV) las
columnas de producto del Excel maestro -- Producto activo, Destacado editorial,
Recomendado, Más vendido, Nuevo, Edición limitada, Etiqueta inicial -- y las
aplica sobre data/menu.json, para poder prender, apagar y etiquetar productos
desde la computadora del local sin correr el build en la PC del dueño.

Mismo circuito que disponibilidad, calendario y precios: URL https de
docs.google.com en la repo variable PRODUCTOS_CSV_URL (no es un secreto),
descarga restringida a Google, mensajes que nunca reproducen el contenido de
una celda.

Qué hace cada columna (es lo mismo que ya hace el Excel, ver
extract_common.badges_for):
  activo               No = el producto se oculta del menú (también su precio).
  destacado            carrusel de destacados, panel grande de su categoría y
                       etiqueta "Favorito de la casa".
  recomendado          etiqueta "Recomendado de la casa".
  mas_vendido          etiqueta "Lo más pedido".
  nuevo                etiqueta "Nuevo".
  (Solo se muestra UNA de las cuatro etiquetas por producto, en este orden de
  prioridad: destacado, recomendado, más vendido, nuevo.)
  edicion_limitada     el sitio todavía NO la usa (igual que en el Excel): se
                       guarda en la hoja pero no cambia nada de lo que se ve.
  etiqueta_inicial     texto corto que aparece en la ficha del producto (solo
                       en español).

Lo que la hoja NO puede hacer: publicar un producto que está en pausa en el
Excel. Su descripción y sus traducciones a 4 idiomas se publican desde el Excel
maestro (y la repo es pública: esos textos no viajan por la hoja). Si en la hoja
se pone Sí a un producto que todavía no está publicado, se informa, no se
inventa nada, y el resto de la hoja se aplica igual.

A diferencia de los precios, acá no es todo o nada: cada fila es independiente
y un error en una no frena las demás (lo peor que puede pasar es una etiqueta
equivocada, y se ve y se arregla). Un error estructural (encabezado, demasiadas
filas, demasiados productos ocultos de una vez) sí frena todo. Cualquier
problema marca productos_ok=false y el workflow pone la corrida en rojo
DESPUÉS de publicar (job avisar-hojas).

Columnas (primera fila = encabezado, en cualquier orden):
  id, activo, destacado, recomendado, mas_vendido, nuevo, edicion_limitada,
  etiqueta_inicial                    obligatorias
  nombre, categoria, estado, nota     opcionales, solo para las personas
Los valores Sí/No pueden ir con o sin tilde y en cualquier mayúscula; en las
columnas de etiquetas, vacío = No. "activo" no puede quedar vacío.

    python build/aplicar_productos.py                       # CI: PRODUCTOS_CSV_URL
    python build/aplicar_productos.py --comparar            # solo lectura
    python build/aplicar_productos.py --plantilla-desde-excel [ruta.xlsx]

Solo biblioteca estándar (openpyxl solo para --plantilla-desde-excel).
"""
import csv
import io
import json
import os
import re
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

import extract_common as ec  # noqa: E402
import hoja_publica  # noqa: E402

MENU_JSON = HERE.parent / "data" / "menu.json"

ENCABEZADOS_PERMITIDOS = ("id", "nombre", "categoria", "estado", "activo", "destacado", "recomendado",
                          "mas_vendido", "nuevo", "edicion_limitada", "etiqueta_inicial", "nota")
ENCABEZADOS_OBLIGATORIOS = ("id", "activo", "destacado", "recomendado", "mas_vendido", "nuevo",
                            "edicion_limitada", "etiqueta_inicial")
COLUMNAS_SI_NO = ("destacado", "recomendado", "mas_vendido", "nuevo", "edicion_limitada")
MAX_FILAS = 500
MAX_OCULTOS = 15  # más que esto de una vez se considera un error de la hoja; se hace desde el Excel
MAX_ETIQUETA = 40
_RE_ID = re.compile(r"^[A-Z0-9]{3,10}$")


class ProductosInvalidos(Exception):
    """Hoja de productos que no se puede confiar (error estructural). El mensaje
    nunca incluye la URL ni el texto crudo de una celda."""


_RedirectHandlerRestringido = hoja_publica.manejador_de_redirecciones(ProductosInvalidos, "de productos")


def _descargar(url: str, timeout: int = 10) -> str:
    return hoja_publica.descargar_csv(url, "de productos", ProductosInvalidos, timeout)


def _si_no(texto: str):
    """'Sí'/'si'/'SI' -> True; 'No'/'' -> False; cualquier otra cosa -> None."""
    t = texto.strip().lower()
    if t in ("sí", "si"):
        return True
    if t in ("no", ""):
        return False
    return None


def etiqueta_valida(texto: str) -> bool:
    """Texto corto y sin nada raro. Se muestra con textContent, pero igual se
    rechazan etiquetas HTML y enlaces."""
    if not texto or len(texto) > MAX_ETIQUETA or texto != texto.strip() or not texto.isprintable():
        return False
    bajo = texto.lower()
    return not any(c in texto for c in "<>") and "://" not in bajo and "www." not in bajo


# ----------------------------------------------------------------------------
# Lectura y validación de la hoja
# ----------------------------------------------------------------------------
def parsear_csv(contenido: str, data: dict) -> tuple:
    """Devuelve (filas, problemas). `filas` = {id: {activo, destacado, recomendado,
    mas_vendido, nuevo, edicion_limitada, etiqueta}} con las filas VÁLIDAS (cada
    una independiente). Si `filas` es None hay un error estructural y no se
    aplica nada. Nunca levanta."""
    registros = list(csv.reader(io.StringIO(contenido)))
    if not registros:
        return None, ["La hoja de productos está vacía (ni encabezado)."]
    encabezado = [c.strip().lower() for c in registros[0]]
    estructurales = []
    desconocidos = [c for c in encabezado if c and c not in ENCABEZADOS_PERMITIDOS]
    if desconocidos:
        estructurales.append(f"El encabezado tiene {len(desconocidos)} columna(s) que no existen. "
                             f"Solo se permiten: {', '.join(ENCABEZADOS_PERMITIDOS)}.")
    no_vacios = [c for c in encabezado if c]
    if len(set(no_vacios)) != len(no_vacios):
        estructurales.append("El encabezado tiene columnas repetidas.")
    faltan = [c for c in ENCABEZADOS_OBLIGATORIOS if c not in encabezado]
    if faltan:
        estructurales.append(f"Faltan columnas obligatorias en el encabezado: {', '.join(faltan)}.")
    if len(registros) - 1 > MAX_FILAS:
        estructurales.append(f"Demasiadas filas ({len(registros) - 1}); el máximo es {MAX_FILAS}.")
    if estructurales:
        return None, estructurales

    indice = {c: i for i, c in enumerate(encabezado) if c}

    def celda(fila, nombre):
        i = indice.get(nombre)
        return fila[i].strip() if i is not None and i < len(fila) else ""

    publicados = {p["id"] for p in data["prods"]}
    filas, problemas, vistos = {}, [], {}
    for numero, fila in enumerate(registros[1:], start=2):
        if not any(c.strip() for c in fila):
            continue
        pid = celda(fila, "id")
        if not _RE_ID.match(pid):
            problemas.append(f"Fila {numero}: el ID no tiene un formato válido.")
            continue
        if pid in vistos:
            problemas.append(f"Fila {numero}: repite el ID de la fila {vistos[pid]}; se ignora.")
            continue
        vistos[pid] = numero
        conocido = pid in publicados
        id_visible = f" ('{pid}')" if conocido else ""

        activo = _si_no(celda(fila, "activo"))
        if activo is None or celda(fila, "activo") == "":
            problemas.append(f"Fila {numero}{id_visible}: 'activo' tiene que ser Sí o No.")
            continue
        flags, error = {}, False
        for col in COLUMNAS_SI_NO:
            valor = _si_no(celda(fila, col))
            if valor is None:
                problemas.append(f"Fila {numero}{id_visible}: '{col}' tiene que estar vacío, Sí o No.")
                error = True
                break
            flags[col] = valor
        if error:
            continue
        etiqueta = celda(fila, "etiqueta_inicial")
        if etiqueta and not etiqueta_valida(etiqueta):
            problemas.append(f"Fila {numero}{id_visible}: 'etiqueta_inicial' tiene que ser un texto de hasta "
                             f"{MAX_ETIQUETA} caracteres, sin < > y sin enlaces.")
            continue

        if not conocido:
            if activo:
                problemas.append(f"Fila {numero}: pusiste Sí en 'activo' pero ese producto todavía no está publicado. "
                                 "Su descripción y sus traducciones se completan y publican desde el Excel maestro.")
            continue  # un producto en pausa con No: nada que aplicar
        filas[pid] = dict(flags, activo=activo, etiqueta=etiqueta)

    ocultos = [p for p, f in filas.items() if not f["activo"]]
    if len(ocultos) > MAX_OCULTOS:
        return None, [f"La hoja oculta {len(ocultos)} productos de una vez (el máximo es {MAX_OCULTOS}). "
                      "Si es a propósito, hacelo desde el Excel maestro; si no, revisá la columna 'activo'."]
    for pid in sorted(publicados - set(vistos)):
        problemas.append(f"El producto publicado '{pid}' no tiene fila en la hoja: se deja como está.")
    return filas, problemas


# ----------------------------------------------------------------------------
# Aplicación sobre data/menu.json
# ----------------------------------------------------------------------------
def aplicar_a_datos(data: dict, filas: dict) -> list:
    """Parchea `data` en el lugar. Devuelve las líneas de lo que cambió."""
    cambios = []
    nuevos = []
    for p in data["prods"]:
        fila = filas.get(p["id"])
        if fila is None:
            nuevos.append(p)
            continue
        if not fila["activo"]:
            cambios.append(f"{p['id']}: oculto")
            data["precios"].pop(p["id"], None)
            continue
        antes = json.dumps(p, sort_keys=True, ensure_ascii=False)
        p["dest"] = fila["destacado"]
        p["b"] = ec.badges_for({"destacado": fila["destacado"], "recomendado": fila["recomendado"],
                                "mas_vendido": fila["mas_vendido"], "nuevo": fila["nuevo"], "cat": p["cat"]})
        if fila["etiqueta"]:
            p["tag"] = fila["etiqueta"]
        else:
            p.pop("tag", None)
        if json.dumps(p, sort_keys=True, ensure_ascii=False) != antes:
            cambios.append(f"{p['id']}: etiquetas actualizadas")
        nuevos.append(p)
    data["prods"] = nuevos
    return cambios


def aplicar_a_archivo(url: str, menu_json_path: Path = MENU_JSON, timeout: int = 10) -> tuple:
    """Devuelve (cambios, problemas). Levanta ProductosInvalidos solo ante un
    error estructural (no se toca nada). Reescribe el archivo solo si algo cambió."""
    ruta = Path(menu_json_path)
    data = json.loads(ruta.read_text(encoding="utf-8"))
    filas, problemas = parsear_csv(_descargar(url, timeout), data)
    if filas is None:
        raise ProductosInvalidos("La hoja de productos tiene datos que no se pueden confiar:\n  - "
                                 + "\n  - ".join(problemas))
    cambios = aplicar_a_datos(data, filas)
    if cambios:
        temporal = ruta.with_suffix(".json.tmp")
        temporal.write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
        temporal.replace(ruta)
    return cambios, problemas


def comparar(url: str, menu_json_path: Path = MENU_JSON) -> tuple:
    """Solo lectura: (cambios, problemas) sin escribir nada."""
    data = json.loads(Path(menu_json_path).read_text(encoding="utf-8"))
    filas, problemas = parsear_csv(_descargar(url), data)
    if filas is None:
        raise ProductosInvalidos("La hoja de productos tiene datos que no se pueden confiar:\n  - "
                                 + "\n  - ".join(problemas))
    return aplicar_a_datos(data, filas), problemas


# ----------------------------------------------------------------------------
# Plantilla inicial (todos los productos del Excel maestro)
# ----------------------------------------------------------------------------
def _txt(v) -> str:
    return "" if v is None else str(v).strip()


def estado_de_producto(publicado: bool, nombres, descripciones, precio, activo_excel: str) -> str:
    if publicado:
        return "Publicado"
    if activo_excel == "Sí":
        return "Activo en el Excel pero no publicado: falta correr el build"
    faltan = sum(1 for n, d in zip(nombres, descripciones) if not (n and d))
    motivos = []
    if faltan:
        motivos.append(f"falta nombre o descripción en {faltan} de 5 idiomas")
    if not precio:
        motivos.append("sin precio")
    return "En pausa: " + ("; ".join(motivos) if motivos else "contenido completo, falta activarlo desde el Excel")


def plantilla_desde_hoja(ws, publicados: set, ultima_fila: int = 400) -> str:
    """CSV para pegar en la pestaña `Productos`: una fila por producto del Excel
    maestro (los publicados y los 'en pausa'), con sus valores de hoy. `ws` es
    una hoja de openpyxl 'Productos' (o cualquier objeto con .cell(row=, column=).value)."""
    C = ec.COL
    salida = io.StringIO()
    w = csv.writer(salida, lineterminator="\n")
    w.writerow(["id", "nombre", "categoria", "estado", "activo", "destacado", "recomendado", "mas_vendido",
                "nuevo", "edicion_limitada", "etiqueta_inicial", "nota"])
    for fila in range(5, ultima_fila):
        pid = _txt(ws.cell(row=fila, column=C["id"]).value)
        if not pid:
            break
        nombres = [_txt(ws.cell(row=fila, column=C["nombre"][l]).value) for l in ec.LANGS]
        descripciones = [_txt(ws.cell(row=fila, column=C["desc"][l]).value) for l in ec.LANGS]
        activo = _txt(ws.cell(row=fila, column=C["activo"]).value)
        precio = _txt(ws.cell(row=fila, column=C["precio_chico"]).value) or _txt(ws.cell(row=fila, column=C["precio_grande"]).value)

        def si_no(clave):
            return "Sí" if _txt(ws.cell(row=fila, column=C[clave]).value) == "Sí" else "No"
        w.writerow([pid, nombres[0], _txt(ws.cell(row=fila, column=C["cat"]).value),
                    estado_de_producto(pid in publicados, nombres, descripciones, precio, activo),
                    "Sí" if activo == "Sí" else "No",
                    si_no("destacado"), si_no("recomendado"), si_no("mas_vendido"), si_no("nuevo"),
                    si_no("edicion_limitada"), _txt(ws.cell(row=fila, column=C["etiqueta_inicial"]).value), ""])
    return salida.getvalue()


# ----------------------------------------------------------------------------
def _salida(clave: str, valor: str) -> None:
    ruta = os.environ.get("GITHUB_OUTPUT")
    if ruta:
        with open(ruta, "a", encoding="utf-8") as f:
            f.write(f"{clave}={valor}\n")


def main(argv=None) -> int:
    argv = list(sys.argv[1:] if argv is None else argv)
    if argv[:1] == ["--plantilla-desde-excel"]:
        import openpyxl  # solo acá: CI no lo necesita
        import config_local
        ruta = Path(argv[1]) if len(argv) > 1 else config_local.resolver_ruta_xlsx(None)
        wb = openpyxl.load_workbook(ruta, data_only=True)
        data = json.loads(MENU_JSON.read_text(encoding="utf-8"))
        sys.stdout.write(plantilla_desde_hoja(wb["Productos"], {p["id"] for p in data["prods"]}))
        return 0
    url = os.environ.get("PRODUCTOS_CSV_URL", "").strip()
    if argv[:1] == ["--comparar"]:
        if not url:
            print("Falta la variable de entorno PRODUCTOS_CSV_URL.")
            return 2
        try:
            cambios, problemas = comparar(url)
        except ProductosInvalidos as e:
            print(e)
            return 1
        print("\n".join(cambios) if cambios else "Sin diferencias: la hoja coincide con data/menu.json.")
        for p in problemas:
            print(f"[AVISO] {p}")
        return 0

    if not url:
        print("Productos: no hay PRODUCTOS_CSV_URL configurada -- se usan los productos de data/menu.json.")
        _salida("productos_ok", "true")
        return 0
    try:
        cambios, problemas = aplicar_a_archivo(url)
    except ProductosInvalidos as e:
        print("::warning::Hoja de productos NO aplicada -- se conserva data/menu.json.")
        print(e)
        _salida("productos_ok", "false")
        return 0
    except Exception as e:  # un bug acá no puede frenar la publicación del menú
        print(f"::warning::Hoja de productos NO aplicada ({type(e).__name__}) -- se conserva data/menu.json.")
        _salida("productos_ok", "false")
        return 0
    print(f"Productos: {len(cambios)} cambio(s) aplicados desde la hoja." if cambios else
          "Productos: la hoja coincide con data/menu.json, sin cambios.")
    for c in cambios:
        print(f"  - {c}")
    if problemas:
        print("::warning::Algunas filas de la hoja de productos NO se aplicaron (el resto sí):")
        for p in problemas:
            print(f"  - {p}")
    _salida("productos_ok", "false" if problemas else "true")
    return 0


if __name__ == "__main__":
    sys.exit(main())
